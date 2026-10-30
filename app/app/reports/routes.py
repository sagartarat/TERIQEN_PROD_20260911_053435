from flask import current_app, render_template, request, jsonify, send_file
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.reports.operational_reports_service import (
    get_invoice_report_dashboard,
    get_invoice_report_rows,
    get_agent_payment_report_dashboard,
    get_agent_payment_report_rows,
    get_sla_report_dashboard,
    get_sla_report_rows,
    get_queue_ageing_report,
    export_queue_ageing_report,
    export_report_excel,
    export_report_pdf,
)


def _route_exists(app, endpoint):
    return endpoint in app.view_functions


def _report_error(exc):
    current_app.logger.exception("REPORT_REQUEST_FAILED")
    return jsonify({
        "success": False,
        "message": "The report could not be generated. Please try again or contact an administrator.",
        "rows": [],
        "data": [],
    }), 500


def _export_excel(title, sheet_name, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    total_cols = 12
    if rows:
        total_cols = max(len(rows[0].keys()), 1)

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_cols,
    )

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F4E78")
    title_cell.alignment = Alignment(horizontal="center")

    if not rows:
        ws.cell(row=3, column=1).value = "No data available"
        ws.cell(row=3, column=1).font = Font(bold=True)
    else:
        headers = list(rows[0].keys())

        header_fill = PatternFill("solid", fgColor="305496")
        header_font = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header.replace("_", " ").title()
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_num, row in enumerate(rows, 4):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = row.get(header)
                cell.alignment = Alignment(vertical="top")
                cell.border = thin_border

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            col_letter = ws.cell(row=3, column=col_idx).column_letter

            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))

            ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def register_reports_routes(app):

    # ============================================================
    # OPERATIONAL REPORTS / QUEUE AGEING REPORT
    # ============================================================

    if not _route_exists(app, "reports_operational_page"):
        @app.route(
            "/reports/operational",
            endpoint="reports_operational_page",
        )
        def reports_operational_page():
            return render_template("operational_reports.html")

    if not _route_exists(app, "api_queue_ageing_report"):
        @app.route(
            "/api/reports/queue-ageing",
            endpoint="api_queue_ageing_report",
        )
        def api_queue_ageing_report():
            try:
                filters = request.args.to_dict()
                rows = get_queue_ageing_report(filters)
                return jsonify({"success": True, "rows": rows, "data": rows})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_queue_ageing_export"):
        @app.route(
            "/api/reports/queue-ageing/export",
            endpoint="api_queue_ageing_export",
        )
        def api_queue_ageing_export():
            try:
                filters = request.args.to_dict()

                try:
                    return export_queue_ageing_report(filters)
                except Exception:
                    rows = get_queue_ageing_report(filters)
                    return _export_excel(
                        title="Queue Ageing Report",
                        sheet_name="Queue Ageing",
                        rows=rows,
                        filename="queue_ageing_report.xlsx",
                    )

            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_queue_ageing_export_excel"):
        @app.route("/api/reports/queue-ageing/export-excel", endpoint="api_queue_ageing_export_excel")
        def api_queue_ageing_export_excel():
            try:
                filters = request.args.to_dict()
                return export_report_excel("queue-ageing", filters)
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_queue_ageing_export_pdf"):
        @app.route("/api/reports/queue-ageing/export-pdf", endpoint="api_queue_ageing_export_pdf")
        def api_queue_ageing_export_pdf():
            try:
                filters = request.args.to_dict()
                return export_report_pdf("queue-ageing", filters)
            except Exception as e:
                return _report_error(e)

    # ============================================================
    # INVOICE MANAGEMENT REPORTS
    # ============================================================

    if not _route_exists(app, "reports_invoice_management_page"):
        @app.route(
            "/reports/invoice-management",
            endpoint="reports_invoice_management_page",
        )
        def reports_invoice_management_page():
            return render_template("invoice_management_reports.html")

    if not _route_exists(app, "api_invoice_management_dashboard"):
        @app.route(
            "/api/reports/invoice-management/dashboard",
            endpoint="api_invoice_management_dashboard",
        )
        def api_invoice_management_dashboard():
            try:
                filters = request.args.to_dict()
                data = get_invoice_report_dashboard(filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_invoice_management_rows"):
        @app.route(
            "/api/reports/invoice-management/rows",
            endpoint="api_invoice_management_rows",
        )
        def api_invoice_management_rows():
            try:
                filters = request.args.to_dict()
                report_type = filters.get("report_type", "invoice_detail")
                data = get_invoice_report_rows(report_type, filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_invoice_management_export"):
        @app.route(
            "/api/reports/invoice-management/export",
            endpoint="api_invoice_management_export",
        )
        def api_invoice_management_export():
            try:
                filters = request.args.to_dict()
                report_type = filters.get("report_type", "invoice_detail")
                rows = get_invoice_report_rows(report_type, filters)

                return _export_excel(
                    title="Invoice Management Report",
                    sheet_name="Invoice Report",
                    rows=rows,
                    filename=f"{report_type}.xlsx",
                )

            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_invoice_management_export_excel"):
        @app.route("/api/reports/invoice-management/export-excel", endpoint="api_invoice_management_export_excel")
        def api_invoice_management_export_excel():
            try:
                filters = request.args.to_dict()
                return export_report_excel("invoice-management", filters)
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_invoice_management_export_pdf"):
        @app.route("/api/reports/invoice-management/export-pdf", endpoint="api_invoice_management_export_pdf")
        def api_invoice_management_export_pdf():
            try:
                filters = request.args.to_dict()
                return export_report_pdf("invoice-management", filters)
            except Exception as e:
                return _report_error(e)

    # ============================================================
    # AGENT PAYMENT REPORTS
    # ============================================================

    if not _route_exists(app, "reports_agent_payment_page"):
        @app.route(
            "/reports/agent-payment",
            endpoint="reports_agent_payment_page",
        )
        def reports_agent_payment_page():
            return render_template("agent_payment_reports.html")

    if not _route_exists(app, "agent_payment_reports"):
        app.add_url_rule(
            "/reports/agent-payment",
            endpoint="agent_payment_reports",
            view_func=reports_agent_payment_page,
        )

    if not _route_exists(app, "api_agent_payment_report_dashboard"):
        @app.route(
            "/api/reports/agent-payment/dashboard",
            endpoint="api_agent_payment_report_dashboard",
        )
        def api_agent_payment_report_dashboard():
            try:
                filters = request.args.to_dict()
                data = get_agent_payment_report_dashboard(filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_agent_payment_report_rows"):
        @app.route(
            "/api/reports/agent-payment/rows",
            endpoint="api_agent_payment_report_rows",
        )
        def api_agent_payment_report_rows():
            try:
                filters = request.args.to_dict()
                report_type = filters.get("report_type", "payment_detail")
                data = get_agent_payment_report_rows(report_type, filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_agent_payment_report_export"):
        @app.route(
            "/api/reports/agent-payment/export",
            endpoint="api_agent_payment_report_export",
        )
        def api_agent_payment_report_export():
            try:
                filters = request.args.to_dict()
                report_type = filters.get("report_type", "payment_detail")
                rows = get_agent_payment_report_rows(report_type, filters)

                return _export_excel(
                    title="Agent Payment Report",
                    sheet_name="Agent Payment",
                    rows=rows,
                    filename=f"agent_payment_{report_type}.xlsx",
                )

            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_agent_payment_report_export_excel"):
        @app.route("/api/reports/agent-payment/export-excel", endpoint="api_agent_payment_report_export_excel")
        def api_agent_payment_report_export_excel():
            try:
                filters = request.args.to_dict()
                return export_report_excel("agent-payment", filters)
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_agent_payment_report_export_pdf"):
        @app.route("/api/reports/agent-payment/export-pdf", endpoint="api_agent_payment_report_export_pdf")
        def api_agent_payment_report_export_pdf():
            try:
                filters = request.args.to_dict()
                return export_report_pdf("agent-payment", filters)
            except Exception as e:
                return _report_error(e)

    # ============================================================
    # SLA REPORTS
    # ============================================================

    if not _route_exists(app, "reports_sla_page"):
        @app.route(
            "/reports/sla",
            endpoint="reports_sla_page",
        )
        def reports_sla_page():
            return render_template("sla_reports.html")

    if not _route_exists(app, "api_sla_report_dashboard"):
        @app.route(
            "/api/reports/sla/dashboard",
            endpoint="api_sla_report_dashboard",
        )
        def api_sla_report_dashboard():
            try:
                filters = request.args.to_dict()
                data = get_sla_report_dashboard(filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_sla_report_rows"):
        @app.route(
            "/api/reports/sla/rows",
            endpoint="api_sla_report_rows",
        )
        def api_sla_report_rows():
            try:
                filters = request.args.to_dict()
                data = get_sla_report_rows(filters)
                return jsonify({"success": True, "data": data})
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_sla_report_export"):
        @app.route(
            "/api/reports/sla/export",
            endpoint="api_sla_report_export",
        )
        def api_sla_report_export():
            try:
                filters = request.args.to_dict()
                rows = get_sla_report_rows(filters)

                return _export_excel(
                    title="SLA Report",
                    sheet_name="SLA Report",
                    rows=rows,
                    filename="sla_report.xlsx",
                )

            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_sla_report_export_excel"):
        @app.route("/api/reports/sla/export-excel", endpoint="api_sla_report_export_excel")
        def api_sla_report_export_excel():
            try:
                filters = request.args.to_dict()
                return export_report_excel("sla", filters)
            except Exception as e:
                return _report_error(e)

    if not _route_exists(app, "api_sla_report_export_pdf"):
        @app.route("/api/reports/sla/export-pdf", endpoint="api_sla_report_export_pdf")
        def api_sla_report_export_pdf():
            try:
                filters = request.args.to_dict()
                return export_report_pdf("sla", filters)
            except Exception as e:
                return _report_error(e)