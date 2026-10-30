from decimal import Decimal
from datetime import date, datetime

try:
    from app.common.db import get_db_connection
except Exception:
    get_db_connection = None

try:
    from app.common.db import get_db
except Exception:
    get_db = None


_TABLE_COLUMNS_CACHE = {}


def _get_conn():
    if get_db_connection:
        return get_db_connection()
    if get_db:
        return get_db()
    raise Exception("Database connection function not found in app.common.db")


def _to_json_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def _normalise_rows(rows):
    clean_rows = []
    for row in rows:
        clean_rows.append({k: _to_json_value(v) for k, v in dict(row).items()})
    return clean_rows


def _fetch_all(sql, params=None):
    conn = _get_conn()
    cursor = None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(sql, params or [])
        rows = cursor.fetchall()
        return _normalise_rows(rows)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def _fetch_one(sql, params=None):
    rows = _fetch_all(sql, params)
    return rows[0] if rows else {}


def _get_table_columns(table_name):
    if table_name in _TABLE_COLUMNS_CACHE:
        return _TABLE_COLUMNS_CACHE[table_name]

    conn = _get_conn()
    cursor = None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(f"SHOW COLUMNS FROM {table_name}")
        rows = cursor.fetchall()
        cols = {r["Field"] for r in rows}
        _TABLE_COLUMNS_CACHE[table_name] = cols
        return cols
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def _first_existing(columns, candidates):
    for col in candidates:
        if col in columns:
            return col
    return None


def _col(alias, column, default="NULL"):
    if column:
        return f"{alias}.{column}"
    return default


# ============================================================
# INVOICE MANAGEMENT REPORTS
# ============================================================

def _build_invoice_where(filters, alias="im"):
    where = ["1=1"]
    params = []

    from_date = filters.get("from_date")
    to_date = filters.get("to_date")
    hospital_name = filters.get("hospital_name")
    invoice_number = filters.get("invoice_number")
    print_status = filters.get("print_status")

    invoice_date_expr = f"DATE(COALESCE({alias}.status_date, {alias}.inventered_at))"

    if from_date:
        where.append(f"{invoice_date_expr} >= %s")
        params.append(from_date)

    if to_date:
        where.append(f"{invoice_date_expr} <= %s")
        params.append(to_date)

    if hospital_name:
        where.append(f"{alias}.hospital_name LIKE %s")
        params.append(f"%{hospital_name}%")

    if invoice_number:
        where.append(f"{alias}.invoice_number LIKE %s")
        params.append(f"%{invoice_number}%")

    if print_status and print_status.upper() != "ALL":
        where.append(f"UPPER(COALESCE({alias}.print_status,'')) = %s")
        params.append(print_status.upper())

    return " AND ".join(where), params


def _build_recon_where(filters, alias="ir"):
    where = ["1=1"]
    params = []

    from_date = filters.get("from_date")
    to_date = filters.get("to_date")
    hospital_name = filters.get("hospital_name")
    invoice_number = filters.get("invoice_number")
    recon_status = filters.get("recon_status")

    recon_date_expr = f"DATE(COALESCE({alias}.pay_date, {alias}.status_date))"

    if from_date:
        where.append(f"{recon_date_expr} >= %s")
        params.append(from_date)

    if to_date:
        where.append(f"{recon_date_expr} <= %s")
        params.append(to_date)

    if hospital_name:
        where.append(f"{alias}.hospital_name LIKE %s")
        params.append(f"%{hospital_name}%")

    if invoice_number:
        where.append(f"{alias}.invoice_number LIKE %s")
        params.append(f"%{invoice_number}%")

    if recon_status and recon_status.upper() != "ALL":
        where.append(f"UPPER(COALESCE({alias}.recon_status,'')) = %s")
        params.append(recon_status.upper())

    return " AND ".join(where), params


def get_invoice_report_dashboard(filters):
    invoice_where, invoice_params = _build_invoice_where(filters, "im")

    invoice_sql = f"""
        SELECT
            COUNT(*) AS invoices_generated,
            SUM(CASE WHEN UPPER(COALESCE(im.print_status,'')) IN ('PRINTED','PRINT') THEN 1 ELSE 0 END) AS printed_invoices,
            SUM(CASE WHEN UPPER(COALESCE(im.print_status,'')) NOT IN ('PRINTED','PRINT') THEN 1 ELSE 0 END) AS pending_print,
            COALESCE(SUM(im.inv_amount),0) AS invoice_amount,
            COALESCE(SUM(im.totalgst_amount),0) AS gst_amount,
            COALESCE(SUM(im.totaltds_amount),0) AS tds_amount,
            COALESCE(SUM(im.netinv_amount),0) AS net_invoice_amount,
            COALESCE(SUM(pay.received_amount),0) AS received_amount,
            SUM(CASE WHEN UPPER(COALESCE(pay.latest_recon_status,'')) = 'MATCHED' THEN 1 ELSE 0 END) AS paid_invoices,
            SUM(CASE WHEN UPPER(COALESCE(pay.latest_recon_status,'')) = 'PARTIAL MATCH' THEN 1 ELSE 0 END) AS partially_paid_invoices
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount,
                MAX(recon_status) AS latest_recon_status
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {invoice_where}
    """

    invoice_data = _fetch_one(invoice_sql, invoice_params)

    net_invoice_amount = float(invoice_data.get("net_invoice_amount") or 0)
    received_amount = float(invoice_data.get("received_amount") or 0)
    outstanding_amount = max(net_invoice_amount - received_amount, 0)

    collection_ratio = 0
    outstanding_ratio = 0

    if net_invoice_amount > 0:
        collection_ratio = round((received_amount / net_invoice_amount) * 100, 2)
        outstanding_ratio = round((outstanding_amount / net_invoice_amount) * 100, 2)

    monthly_sql = f"""
        SELECT
            DATE_FORMAT(COALESCE(im.status_date, im.inventered_at), '%Y-%m') AS month_name,
            COUNT(*) AS invoice_count,
            COALESCE(SUM(im.netinv_amount),0) AS invoice_value,
            COALESCE(SUM(pay.received_amount),0) AS received_amount,
            COALESCE(SUM(im.netinv_amount),0) - COALESCE(SUM(pay.received_amount),0) AS outstanding_amount
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {invoice_where}
        GROUP BY DATE_FORMAT(COALESCE(im.status_date, im.inventered_at), '%Y-%m')
        ORDER BY month_name
    """

    hospital_sql = f"""
        SELECT
            im.hospital_name,
            COUNT(*) AS invoice_count,
            COALESCE(SUM(im.netinv_amount),0) AS invoice_value,
            COALESCE(SUM(pay.received_amount),0) AS received_amount,
            COALESCE(SUM(im.netinv_amount),0) - COALESCE(SUM(pay.received_amount),0) AS outstanding_amount
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {invoice_where}
        GROUP BY im.hospital_name
        ORDER BY outstanding_amount DESC
        LIMIT 10
    """

    return {
        "summary": {
            **invoice_data,
            "outstanding_amount": round(outstanding_amount, 2),
            "collection_ratio": collection_ratio,
            "outstanding_ratio": outstanding_ratio,
        },
        "monthly_trend": _fetch_all(monthly_sql, invoice_params),
        "hospital_outstanding": _fetch_all(hospital_sql, invoice_params),
    }


def get_invoice_report_rows(report_type, filters):
    if report_type == "collection_report":
        return _collection_report(filters)

    if report_type == "hospital_outstanding":
        return _hospital_outstanding_report(filters)

    if report_type == "invoice_ageing":
        return _invoice_ageing_report(filters)

    if report_type == "monthly_revenue":
        return _monthly_revenue_report(filters)

    return _invoice_detail_report(filters)


def _invoice_detail_report(filters):
    where, params = _build_invoice_where(filters, "im")

    sql = f"""
        SELECT
            im.invoice_number,
            im.hospital_name,
            im.patient_name,
            im.patient_passport,
            DATE(COALESCE(im.status_date, im.inventered_at)) AS invoice_date,
            im.wrkflw_status,
            im.print_status,
            COALESCE(im.inv_amount,0) AS invoice_amount,
            COALESCE(im.totalgst_amount,0) AS gst_amount,
            COALESCE(im.totaltds_amount,0) AS tds_amount,
            COALESCE(im.netinv_amount,0) AS net_invoice_amount,
            COALESCE(pay.received_amount,0) AS received_amount,
            COALESCE(im.netinv_amount,0) - COALESCE(pay.received_amount,0) AS outstanding_amount,
            CASE
                WHEN COALESCE(pay.received_amount,0) >= COALESCE(im.netinv_amount,0)
                     AND COALESCE(im.netinv_amount,0) > 0 THEN 'PAID'
                WHEN COALESCE(pay.received_amount,0) > 0 THEN 'PARTIAL PAID'
                ELSE 'PENDING'
            END AS payment_status
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {where}
        ORDER BY DATE(COALESCE(im.status_date, im.inventered_at)) DESC, im.invoice_number DESC
    """

    return _fetch_all(sql, params)


def _collection_report(filters):
    where, params = _build_recon_where(filters, "ir")

    sql = f"""
        SELECT
            ir.invoice_number,
            ir.hospital_name,
            ir.pay_date,
            ir.status_date,
            COALESCE(ir.netinv_amount,0) AS net_invoice_amount,
            COALESCE(ir.recvd_amount,0) AS received_amount,
            COALESCE(ir.diff_amount,0) AS difference_amount,
            ir.utrref_no,
            ir.recon_status,
            ir.wrkflw_status,
            ir.recon_remarks
        FROM invrecon_master ir
        WHERE {where}
        ORDER BY COALESCE(ir.pay_date, ir.status_date) DESC, ir.invoice_number DESC
    """

    return _fetch_all(sql, params)


def _hospital_outstanding_report(filters):
    where, params = _build_invoice_where(filters, "im")

    sql = f"""
        SELECT
            im.hospital_name,
            COUNT(*) AS invoice_count,
            COALESCE(SUM(im.netinv_amount),0) AS net_invoice_amount,
            COALESCE(SUM(pay.received_amount),0) AS received_amount,
            COALESCE(SUM(im.netinv_amount),0) - COALESCE(SUM(pay.received_amount),0) AS outstanding_amount,
            ROUND(
                CASE
                    WHEN COALESCE(SUM(im.netinv_amount),0) = 0 THEN 0
                    ELSE COALESCE(SUM(pay.received_amount),0) / COALESCE(SUM(im.netinv_amount),0) * 100
                END, 2
            ) AS collection_ratio_percent
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {where}
        GROUP BY im.hospital_name
        ORDER BY outstanding_amount DESC
    """

    return _fetch_all(sql, params)


def _invoice_ageing_report(filters):
    where, params = _build_invoice_where(filters, "im")

    sql = f"""
        SELECT
            im.invoice_number,
            im.hospital_name,
            im.patient_name,
            DATE(COALESCE(im.status_date, im.inventered_at)) AS invoice_date,
            DATEDIFF(CURDATE(), DATE(COALESCE(im.status_date, im.inventered_at))) AS ageing_days,
            im.print_status,
            COALESCE(im.netinv_amount,0) AS net_invoice_amount,
            COALESCE(pay.received_amount,0) AS received_amount,
            COALESCE(im.netinv_amount,0) - COALESCE(pay.received_amount,0) AS outstanding_amount,
            CASE
                WHEN COALESCE(im.netinv_amount,0) - COALESCE(pay.received_amount,0) <= 0 THEN 'PAID'
                WHEN DATEDIFF(CURDATE(), DATE(COALESCE(im.status_date, im.inventered_at))) <= 7 THEN '0-7 DAYS'
                WHEN DATEDIFF(CURDATE(), DATE(COALESCE(im.status_date, im.inventered_at))) <= 15 THEN '8-15 DAYS'
                WHEN DATEDIFF(CURDATE(), DATE(COALESCE(im.status_date, im.inventered_at))) <= 30 THEN '16-30 DAYS'
                ELSE '30+ DAYS'
            END AS ageing_bucket
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {where}
        ORDER BY outstanding_amount DESC, ageing_days DESC
    """

    return _fetch_all(sql, params)


def _monthly_revenue_report(filters):
    where, params = _build_invoice_where(filters, "im")

    sql = f"""
        SELECT
            DATE_FORMAT(COALESCE(im.status_date, im.inventered_at), '%Y-%m') AS month_name,
            COUNT(*) AS invoice_count,
            COALESCE(SUM(im.inv_amount),0) AS invoice_amount,
            COALESCE(SUM(im.totalgst_amount),0) AS gst_amount,
            COALESCE(SUM(im.totaltds_amount),0) AS tds_amount,
            COALESCE(SUM(im.netinv_amount),0) AS net_invoice_amount,
            COALESCE(SUM(pay.received_amount),0) AS received_amount,
            COALESCE(SUM(im.netinv_amount),0) - COALESCE(SUM(pay.received_amount),0) AS outstanding_amount
        FROM invoice_master im
        LEFT JOIN (
            SELECT
                invoice_number,
                SUM(COALESCE(recvd_amount,0)) AS received_amount
            FROM invrecon_master
            GROUP BY invoice_number
        ) pay ON pay.invoice_number = im.invoice_number
        WHERE {where}
        GROUP BY DATE_FORMAT(COALESCE(im.status_date, im.inventered_at), '%Y-%m')
        ORDER BY month_name
    """

    return _fetch_all(sql, params)


# ============================================================
# SLA REPORTS
# ============================================================

def _sla_meta():
    osl_cols = _get_table_columns("operation_queue_sla_log")
    oqm_cols = _get_table_columns("operation_queue_master")

    return {
        "osl_cols": osl_cols,
        "oqm_cols": oqm_cols,
        "status_col": _first_existing(osl_cols, ["status", "sla_status"]),
        "elapsed_col": _first_existing(osl_cols, ["minutes_elapsed", "elapsed_minutes", "elapsed_mins"]),
        "target_col": _first_existing(osl_cols, ["target_minutes", "target_duration", "target_duration_mins"]),
        "near_col": _first_existing(osl_cols, ["near_breach_minutes", "near_breach_mins"]),
        "esc_col": _first_existing(osl_cols, ["current_escalation_level", "escalation_level"]),
        "rule_code_col": _first_existing(osl_cols, ["rule_code", "sla_rule_code"]),
        "rule_name_col": _first_existing(osl_cols, ["rule_name", "sla_rule_name"]),
        "started_col": _first_existing(osl_cols, ["started_on", "start_time", "sla_started_on"]),
        "stopped_col": _first_existing(osl_cols, ["stopped_on", "stop_time", "sla_stopped_on"]),
        "created_col": _first_existing(osl_cols, ["created_on", "created_at"]),
        "updated_col": _first_existing(osl_cols, ["updated_on", "updated_at"]),
        "operation_queue_id_col": _first_existing(osl_cols, ["operation_queue_id", "queue_id"]),
        "oqm_id_col": _first_existing(oqm_cols, ["id", "operation_queue_id"]),
    }


def _sla_join_clause(meta):
    if meta["operation_queue_id_col"] and meta["oqm_id_col"]:
        return f"""
            LEFT JOIN operation_queue_master oqm
                ON oqm.{meta["oqm_id_col"]} = osl.{meta["operation_queue_id_col"]}
        """

    if "channel_thread_id" in meta["osl_cols"] and "channel_thread_id" in meta["oqm_cols"]:
        return """
            LEFT JOIN operation_queue_master oqm
                ON oqm.channel_thread_id = osl.channel_thread_id
        """

    return """
        LEFT JOIN operation_queue_master oqm
            ON 1 = 0
    """


def _build_sla_where(filters):
    meta = _sla_meta()
    where = ["1=1"]
    params = []

    date_col = meta["created_col"] or meta["started_col"]
    date_expr = f"DATE(osl.{date_col})" if date_col else "CURDATE()"

    status_expr = _col("osl", meta["status_col"], "''")
    priority_expr = "oqm.priority_status" if "priority_status" in meta["oqm_cols"] else "''"
    hospital_expr = "oqm.hospital_shrtnm" if "hospital_shrtnm" in meta["oqm_cols"] else "''"
    agency_expr = "oqm.agent_shrtnm" if "agent_shrtnm" in meta["oqm_cols"] else "''"

    if filters.get("from_date"):
        where.append(f"{date_expr} >= %s")
        params.append(filters.get("from_date"))

    if filters.get("to_date"):
        where.append(f"{date_expr} <= %s")
        params.append(filters.get("to_date"))

    if filters.get("sla_status"):
        where.append(f"UPPER(COALESCE({status_expr},'')) = %s")
        params.append(filters.get("sla_status").upper())

    if filters.get("priority_status") and "priority_status" in meta["oqm_cols"]:
        where.append(f"UPPER(COALESCE({priority_expr},'')) = %s")
        params.append(filters.get("priority_status").upper())

    if filters.get("hospital_shrtnm") and "hospital_shrtnm" in meta["oqm_cols"]:
        where.append(f"{hospital_expr} LIKE %s")
        params.append(f"%{filters.get('hospital_shrtnm')}%")

    if filters.get("agent_shrtnm") and "agent_shrtnm" in meta["oqm_cols"]:
        where.append(f"{agency_expr} LIKE %s")
        params.append(f"%{filters.get('agent_shrtnm')}%")

    if filters.get("search"):
        like = f"%{filters.get('search')}%"
        search_parts = []

        for col in [
            "channel_thread_id",
            "patpassport_no",
            "patgiven_nm",
            "patsur_nm",
            "hospital_shrtnm",
            "agent_shrtnm",
            "patient_country",
        ]:
            if col in meta["oqm_cols"]:
                search_parts.append(f"oqm.{col} LIKE %s")
                params.append(like)

        if search_parts:
            where.append("(" + " OR ".join(search_parts) + ")")

    return " AND ".join(where), params, meta


def get_sla_report_dashboard(filters):
    where, params, meta = _build_sla_where(filters)

    join_clause = _sla_join_clause(meta)

    status_expr = _col("osl", meta["status_col"], "''")
    elapsed_expr = f"COALESCE({_col('osl', meta['elapsed_col'], '0')},0)"
    esc_expr = f"COALESCE({_col('osl', meta['esc_col'], '0')},0)"

    sql = f"""
        SELECT
            COUNT(*) AS total_records,
            SUM(CASE WHEN UPPER(COALESCE({status_expr},'')) IN ('MET','WITHIN SLA','OPEN','ACTIVE') THEN 1 ELSE 0 END) AS within_sla,
            SUM(CASE WHEN UPPER(COALESCE({status_expr},'')) = 'NEAR BREACH' THEN 1 ELSE 0 END) AS near_breach,
            SUM(CASE WHEN UPPER(COALESCE({status_expr},'')) = 'BREACHED' THEN 1 ELSE 0 END) AS breached,
            SUM(CASE WHEN {esc_expr} > 0 THEN 1 ELSE 0 END) AS escalated,
            ROUND(AVG({elapsed_expr}),2) AS avg_elapsed_minutes
        FROM operation_queue_sla_log osl
        {join_clause}
        WHERE {where}
    """

    filters_data = {
        "hospitals": [],
        "agencies": [],
        "statuses": [],
        "priorities": [],
    }

    if "hospital_shrtnm" in meta["oqm_cols"]:
        filters_data["hospitals"] = [r["value"] for r in _fetch_all("""
            SELECT DISTINCT hospital_shrtnm AS value
            FROM operation_queue_master
            WHERE hospital_shrtnm IS NOT NULL AND hospital_shrtnm <> ''
            ORDER BY hospital_shrtnm
        """)]

    if "agent_shrtnm" in meta["oqm_cols"]:
        filters_data["agencies"] = [r["value"] for r in _fetch_all("""
            SELECT DISTINCT agent_shrtnm AS value
            FROM operation_queue_master
            WHERE agent_shrtnm IS NOT NULL AND agent_shrtnm <> ''
            ORDER BY agent_shrtnm
        """)]

    if meta["status_col"]:
        filters_data["statuses"] = [r["value"] for r in _fetch_all(f"""
            SELECT DISTINCT {meta["status_col"]} AS value
            FROM operation_queue_sla_log
            WHERE {meta["status_col"]} IS NOT NULL AND {meta["status_col"]} <> ''
            ORDER BY {meta["status_col"]}
        """)]

    if "priority_status" in meta["oqm_cols"]:
        filters_data["priorities"] = [r["value"] for r in _fetch_all("""
            SELECT DISTINCT priority_status AS value
            FROM operation_queue_master
            WHERE priority_status IS NOT NULL AND priority_status <> ''
            ORDER BY priority_status
        """)]

    return {
        "kpis": _fetch_one(sql, params),
        "filters": filters_data,
    }


def get_sla_report_rows(filters):
    where, params, meta = _build_sla_where(filters)

    join_clause = _sla_join_clause(meta)

    status_expr = _col("osl", meta["status_col"], "''")
    elapsed_expr = f"COALESCE({_col('osl', meta['elapsed_col'], '0')},0)"
    esc_expr = f"COALESCE({_col('osl', meta['esc_col'], '0')},0)"

    sql = f"""
        SELECT
            {_col("oqm", "channel_thread_id" if "channel_thread_id" in meta["oqm_cols"] else None, "NULL")} AS channel_thread_id,
            {_col("oqm", "patpassport_no" if "patpassport_no" in meta["oqm_cols"] else None, "NULL")} AS patpassport_no,
            CONCAT(
                COALESCE({_col("oqm", "patgiven_nm" if "patgiven_nm" in meta["oqm_cols"] else None, "''")},''),
                ' ',
                COALESCE({_col("oqm", "patsur_nm" if "patsur_nm" in meta["oqm_cols"] else None, "''")},'')
            ) AS patient_name,
            {_col("oqm", "patient_country" if "patient_country" in meta["oqm_cols"] else None, "NULL")} AS patient_country,
            {_col("oqm", "hospital_shrtnm" if "hospital_shrtnm" in meta["oqm_cols"] else None, "NULL")} AS hospital_shrtnm,
            {_col("oqm", "agent_shrtnm" if "agent_shrtnm" in meta["oqm_cols"] else None, "NULL")} AS agent_shrtnm,
            {_col("oqm", "priority_status" if "priority_status" in meta["oqm_cols"] else None, "NULL")} AS priority_status,
            {_col("oqm", "wrkflw_status" if "wrkflw_status" in meta["oqm_cols"] else None, "NULL")} AS wrkflw_status,
            {_col("osl", meta["rule_code_col"], "NULL")} AS rule_code,
            {_col("osl", meta["rule_name_col"], "NULL")} AS rule_name,
            {status_expr} AS sla_status,
            {elapsed_expr} AS minutes_elapsed,
            {_col("osl", meta["target_col"], "0")} AS target_minutes,
            {_col("osl", meta["near_col"], "0")} AS near_breach_minutes,
            {esc_expr} AS current_escalation_level,
            {_col("osl", meta["started_col"], "NULL")} AS started_on,
            {_col("osl", meta["stopped_col"], "NULL")} AS stopped_on,
            {_col("osl", meta["created_col"], "NULL")} AS created_on,
            {_col("osl", meta["updated_col"], "NULL")} AS updated_on
        FROM operation_queue_sla_log osl
        {join_clause}
        WHERE {where}
        ORDER BY
            CASE
                WHEN UPPER(COALESCE({status_expr},'')) = 'BREACHED' THEN 1
                WHEN UPPER(COALESCE({status_expr},'')) = 'NEAR BREACH' THEN 2
                WHEN {esc_expr} > 0 THEN 3
                ELSE 4
            END,
            {elapsed_expr} DESC
    """

    return _fetch_all(sql, params)

# ============================================================
# QUEUE AGEING REPORT
# ============================================================

def _queue_ageing_meta():
    oqm_cols = _get_table_columns("operation_queue_master")

    return {
        "oqm_cols": oqm_cols,
        "id_col": _first_existing(oqm_cols, ["id", "operation_queue_id"]),
        "date_col": _first_existing(
            oqm_cols,
            [
                "created_on",
                "created_at",
                "received_on",
                "received_at",
                "received_dt",
                "message_received_on",
                "latest_message_date",
                "latest_message_dt",
                "assigned_at",
                "updated_on",
                "updated_at",
            ],
        ),
        "updated_col": _first_existing(oqm_cols, ["updated_on", "updated_at"]),
        "assigned_col": _first_existing(oqm_cols, ["assigned_at", "assigned_on"]),
    }


def _build_queue_ageing_where(filters, alias="oqm"):
    meta = _queue_ageing_meta()
    oqm_cols = meta["oqm_cols"]
    date_col = meta["date_col"]

    where = ["1=1"]
    params = []

    date_expr = f"DATE({alias}.{date_col})" if date_col else "CURDATE()"

    if date_col and filters.get("from_date"):
        where.append(f"{date_expr} >= %s")
        params.append(filters.get("from_date"))

    if date_col and filters.get("to_date"):
        where.append(f"{date_expr} <= %s")
        params.append(filters.get("to_date"))

    if filters.get("wrkflw_status") and filters.get("wrkflw_status").upper() != "ALL" and "wrkflw_status" in oqm_cols:
        where.append(f"UPPER(COALESCE({alias}.wrkflw_status,'')) = %s")
        params.append(filters.get("wrkflw_status").upper())

    if filters.get("priority_status") and filters.get("priority_status").upper() != "ALL" and "priority_status" in oqm_cols:
        where.append(f"UPPER(COALESCE({alias}.priority_status,'')) = %s")
        params.append(filters.get("priority_status").upper())

    if filters.get("hospital_shrtnm") and "hospital_shrtnm" in oqm_cols:
        where.append(f"{alias}.hospital_shrtnm LIKE %s")
        params.append(f"%{filters.get('hospital_shrtnm')}%")

    if filters.get("agent_shrtnm") and "agent_shrtnm" in oqm_cols:
        where.append(f"{alias}.agent_shrtnm LIKE %s")
        params.append(f"%{filters.get('agent_shrtnm')}%")

    if filters.get("assigned_to") and "assigned_to" in oqm_cols:
        where.append(f"{alias}.assigned_to LIKE %s")
        params.append(f"%{filters.get('assigned_to')}%")

    if filters.get("parse_status") and filters.get("parse_status").upper() != "ALL" and "parse_status" in oqm_cols:
        where.append(f"UPPER(COALESCE({alias}.parse_status,'')) = %s")
        params.append(filters.get("parse_status").upper())

    if filters.get("search"):
        like = f"%{filters.get('search')}%"
        search_parts = []

        for col in [
            "channel_thread_id",
            "patpassport_no",
            "patgiven_nm",
            "patsur_nm",
            "patient_country",
            "hospital_shrtnm",
            "hospital_city",
            "hospital_dept",
            "ivac_center",
            "agent_shrtnm",
            "doctor_nm",
            "wrkflw_status",
            "priority_status",
            "assigned_to",
            "parse_status",
            "intent_remarks",
        ]:
            if col in oqm_cols:
                search_parts.append(f"{alias}.{col} LIKE %s")
                params.append(like)

        if search_parts:
            where.append("(" + " OR ".join(search_parts) + ")")

    return " AND ".join(where), params, meta


def get_queue_ageing_report(filters):
    where, params, meta = _build_queue_ageing_where(filters, "oqm")
    oqm_cols = meta["oqm_cols"]

    date_col = meta["date_col"]
    updated_col = meta["updated_col"]
    assigned_col = meta["assigned_col"]

    queue_date_expr = _col("oqm", date_col, "NULL")
    updated_expr = _col("oqm", updated_col, "NULL")
    assigned_expr = _col("oqm", assigned_col, "NULL")

    patient_name_expr = f"""
        TRIM(CONCAT(
            COALESCE({_col("oqm", "patgiven_nm" if "patgiven_nm" in oqm_cols else None, "''")}, ''),
            ' ',
            COALESCE({_col("oqm", "patsur_nm" if "patsur_nm" in oqm_cols else None, "''")}, '')
        ))
    """

    wrkflw_status_expr = _col("oqm", "wrkflw_status" if "wrkflw_status" in oqm_cols else None, "''")
    priority_expr = _col("oqm", "priority_status" if "priority_status" in oqm_cols else None, "''")

    sql = f"""
        SELECT
            {_col("oqm", "id" if "id" in oqm_cols else None, "NULL")} AS operation_queue_id,
            {_col("oqm", "channel_thread_id" if "channel_thread_id" in oqm_cols else None, "NULL")} AS channel_thread_id,
            {_col("oqm", "channel_code" if "channel_code" in oqm_cols else None, "NULL")} AS channel_code,
            {_col("oqm", "channel_message_id" if "channel_message_id" in oqm_cols else None, "NULL")} AS channel_message_id,
            {_col("oqm", "patpassport_no" if "patpassport_no" in oqm_cols else None, "NULL")} AS patpassport_no,
            {patient_name_expr} AS patient_name,
            {_col("oqm", "patient_country" if "patient_country" in oqm_cols else None, "NULL")} AS patient_country,
            {_col("oqm", "hospital_shrtnm" if "hospital_shrtnm" in oqm_cols else None, "NULL")} AS hospital_shrtnm,
            {_col("oqm", "hospital_city" if "hospital_city" in oqm_cols else None, "NULL")} AS hospital_city,
            {_col("oqm", "hospital_dept" if "hospital_dept" in oqm_cols else None, "NULL")} AS hospital_dept,
            {_col("oqm", "ivac_center" if "ivac_center" in oqm_cols else None, "NULL")} AS ivac_center,
            {_col("oqm", "agent_shrtnm" if "agent_shrtnm" in oqm_cols else None, "NULL")} AS agent_shrtnm,
            {_col("oqm", "assigned_to" if "assigned_to" in oqm_cols else None, "NULL")} AS assigned_to,
            {assigned_expr} AS assigned_at,
            {priority_expr} AS priority_status,
            {wrkflw_status_expr} AS wrkflw_status,
            {_col("oqm", "parse_status" if "parse_status" in oqm_cols else None, "NULL")} AS parse_status,
            {_col("oqm", "intent_remarks" if "intent_remarks" in oqm_cols else None, "NULL")} AS request_for,
            {queue_date_expr} AS queue_created_on,
            {updated_expr} AS last_updated_on,
            CASE
                WHEN {queue_date_expr} IS NULL THEN 0
                ELSE DATEDIFF(CURDATE(), DATE({queue_date_expr}))
            END AS ageing_days,
            CASE
                WHEN {queue_date_expr} IS NULL THEN 0
                ELSE TIMESTAMPDIFF(HOUR, {queue_date_expr}, NOW())
            END AS ageing_hours,
            CASE
                WHEN {queue_date_expr} IS NULL THEN 'UNKNOWN'
                WHEN DATEDIFF(CURDATE(), DATE({queue_date_expr})) <= 1 THEN '0-1 DAYS'
                WHEN DATEDIFF(CURDATE(), DATE({queue_date_expr})) <= 3 THEN '2-3 DAYS'
                WHEN DATEDIFF(CURDATE(), DATE({queue_date_expr})) <= 7 THEN '4-7 DAYS'
                WHEN DATEDIFF(CURDATE(), DATE({queue_date_expr})) <= 15 THEN '8-15 DAYS'
                ELSE '15+ DAYS'
            END AS ageing_bucket
        FROM operation_queue_master oqm
        WHERE {where}
        ORDER BY
            CASE
                WHEN UPPER(COALESCE({wrkflw_status_expr},'')) = 'OPEN' THEN 1
                WHEN UPPER(COALESCE({wrkflw_status_expr},'')) = 'ASSIGNED' THEN 2
                WHEN UPPER(COALESCE({wrkflw_status_expr},'')) = 'IN PROGRESS' THEN 3
                WHEN UPPER(COALESCE({wrkflw_status_expr},'')) = 'MANUAL REVIEW' THEN 4
                WHEN UPPER(COALESCE({wrkflw_status_expr},'')) = 'CLOSED' THEN 9
                ELSE 5
            END,
            CASE
                WHEN UPPER(COALESCE({priority_expr},'')) = 'HIGH' THEN 1
                WHEN UPPER(COALESCE({priority_expr},'')) = 'MEDIUM' THEN 2
                WHEN UPPER(COALESCE({priority_expr},'')) = 'NORMAL' THEN 3
                WHEN UPPER(COALESCE({priority_expr},'')) = 'LOW' THEN 4
                ELSE 5
            END,
            ageing_days DESC,
            queue_created_on DESC
    """

    return _fetch_all(sql, params)


def export_queue_ageing_report(filters):
    from io import BytesIO
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    rows = get_queue_ageing_report(filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Queue Ageing"

    total_cols = 12
    if rows:
        total_cols = max(len(rows[0].keys()), 1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = "Queue Ageing Report"
    title_cell.font = Font(size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F4E78")
    title_cell.alignment = Alignment(horizontal="center")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    if not rows:
        ws.cell(row=4, column=1).value = "No data available"
        ws.cell(row=4, column=1).font = Font(bold=True)
    else:
        headers = list(rows[0].keys())

        header_fill = PatternFill("solid", fgColor="305496")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = header.replace("_", " ").title()
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_num, row in enumerate(rows, 5):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = row.get(header)
                cell.alignment = Alignment(vertical="top")
                cell.border = thin_border

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            # Use a real worksheet cell to get the column letter.
            # This avoids openpyxl MergedCell errors caused by the merged report title row.
            col_letter = ws.cell(row=4, column=col_idx).column_letter

            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))

            ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="queue_ageing_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )



# ============================================================
# AGENT PAYMENT REPORTS
# ============================================================

def _agent_payment_meta():
    apm_cols = _get_table_columns("agent_payment_master")
    apd_cols = _get_table_columns("agent_payment_detail")

    try:
        im_cols = _get_table_columns("invoice_master")
    except Exception:
        im_cols = set()

    return {
        "apm_cols": apm_cols,
        "apd_cols": apd_cols,
        "im_cols": im_cols,
        "payment_no_col": _first_existing(apm_cols, ["payment_no"]),
        "payment_date_col": _first_existing(apm_cols, ["payment_date", "entered_on", "created_on"]),
        "agent_code_col": _first_existing(apm_cols, ["agent_code", "agent_shrtnm"]),
        "agent_name_col": _first_existing(apm_cols, ["agent_name"]),
        "payment_mode_col": _first_existing(apm_cols, ["payment_mode"]),
        "payment_amount_col": _first_existing(apm_cols, ["payment_amount"]),
        "tds_amount_col": _first_existing(apm_cols, ["tds_amount"]),
        "net_paid_amount_col": _first_existing(apm_cols, ["net_paid_amount"]),
        "bank_reference_col": _first_existing(apm_cols, ["bank_reference"]),
        "utr_reference_col": _first_existing(apm_cols, ["utr_reference", "utrref_no"]),
        "payment_status_col": _first_existing(apm_cols, ["payment_status", "status"]),
        "currency_code_col": _first_existing(apm_cols, ["currency_code", "currency"]),
        "remarks_col": _first_existing(apm_cols, ["remarks"]),
        "entered_by_col": _first_existing(apm_cols, ["entered_by", "created_by"]),
        "entered_on_col": _first_existing(apm_cols, ["entered_on", "created_on"]),

        "detail_payment_no_col": _first_existing(apd_cols, ["payment_no"]),
        "invoice_number_col": _first_existing(apd_cols, ["invoice_number", "invoice_no"]),
        "invoice_amount_col": _first_existing(apd_cols, ["invoice_amount"]),
        "already_paid_amount_col": _first_existing(apd_cols, ["already_paid_amount"]),
        "payable_amount_col": _first_existing(apd_cols, ["payable_amount", "agency_pay_amount"]),
        "paid_now_amount_col": _first_existing(apd_cols, ["paid_now_amount", "paid_amount"]),
        "balance_amount_col": _first_existing(apd_cols, ["balance_amount", "outstanding_amount"]),
        "detail_remarks_col": _first_existing(apd_cols, ["remarks"]),
        "created_on_col": _first_existing(apd_cols, ["created_on", "created_at"]),

        "im_invoice_number_col": _first_existing(im_cols, ["invoice_number", "invoice_no"]),
        "im_patient_name_col": _first_existing(im_cols, ["patient_name", "pat_name"]),
        "im_patient_passport_col": _first_existing(im_cols, ["patient_passport", "patpassport_no", "passport_no"]),
        "im_currency_col": _first_existing(im_cols, ["inv_currency", "currency_code", "currency"]),
        "im_hospital_col": _first_existing(im_cols, ["hospital_name", "hospital_shrtnm"]),
    }


def _agent_payment_join_clause(meta):
    if meta["detail_payment_no_col"] and meta["payment_no_col"]:
        return f"""
            INNER JOIN agent_payment_master apm
                ON apm.{meta["payment_no_col"]} = apd.{meta["detail_payment_no_col"]}
        """

    return """
        INNER JOIN agent_payment_master apm
            ON 1 = 0
    """


def _agent_payment_invoice_join_clause(meta):
    if meta["invoice_number_col"] and meta["im_invoice_number_col"]:
        return f"""
            LEFT JOIN invoice_master im
                ON UPPER(TRIM(im.{meta["im_invoice_number_col"]})) = UPPER(TRIM(apd.{meta["invoice_number_col"]}))
        """

    return """
        LEFT JOIN invoice_master im
            ON 1 = 0
    """


def _agent_payment_base_exprs(meta):
    payment_date_col = meta["payment_date_col"]
    payment_date_expr = f"DATE(apm.{payment_date_col})" if payment_date_col else "CURDATE()"
    payment_date_select = f"DATE(apm.{payment_date_col})" if payment_date_col else "NULL"

    payment_no_expr = _col("apm", meta["payment_no_col"], "''")
    agent_code_expr = _col("apm", meta["agent_code_col"], "''")
    agent_name_expr = _col("apm", meta["agent_name_col"], agent_code_expr)
    payment_mode_expr = _col("apm", meta["payment_mode_col"], "''")
    payment_status_expr = _col("apm", meta["payment_status_col"], "''")
    currency_expr = _col("apm", meta["currency_code_col"], _col("im", meta["im_currency_col"], "'INR'"))
    utr_expr = _col("apm", meta["utr_reference_col"], "''")
    bank_ref_expr = _col("apm", meta["bank_reference_col"], "''")
    remarks_expr = _col("apm", meta["remarks_col"], "''")
    entered_by_expr = _col("apm", meta["entered_by_col"], "''")
    entered_on_expr = _col("apm", meta["entered_on_col"], "NULL")

    invoice_number_expr = _col("apd", meta["invoice_number_col"], "''")
    invoice_amount_expr = f"COALESCE({_col('apd', meta['invoice_amount_col'], '0')},0)"
    payable_amount_expr = f"COALESCE({_col('apd', meta['payable_amount_col'], '0')},0)"
    paid_now_amount_expr = f"COALESCE({_col('apd', meta['paid_now_amount_col'], '0')},0)"
    balance_amount_expr = f"COALESCE({_col('apd', meta['balance_amount_col'], '0')},0)"
    already_paid_amount_expr = f"COALESCE({_col('apd', meta['already_paid_amount_col'], '0')},0)"
    payment_amount_expr = f"COALESCE({_col('apm', meta['payment_amount_col'], '0')},0)"
    tds_amount_expr = f"COALESCE({_col('apm', meta['tds_amount_col'], '0')},0)"
    net_paid_amount_expr = f"COALESCE({_col('apm', meta['net_paid_amount_col'], '0')},0)"

    patient_expr = _col("im", meta["im_patient_name_col"], "''")
    passport_expr = _col("im", meta["im_patient_passport_col"], "''")
    hospital_expr = _col("im", meta["im_hospital_col"], "''")

    return {
        "payment_date_expr": payment_date_expr,
        "payment_date_select": payment_date_select,
        "payment_no_expr": payment_no_expr,
        "agent_code_expr": agent_code_expr,
        "agent_name_expr": agent_name_expr,
        "payment_mode_expr": payment_mode_expr,
        "payment_status_expr": payment_status_expr,
        "currency_expr": currency_expr,
        "utr_expr": utr_expr,
        "bank_ref_expr": bank_ref_expr,
        "remarks_expr": remarks_expr,
        "entered_by_expr": entered_by_expr,
        "entered_on_expr": entered_on_expr,
        "invoice_number_expr": invoice_number_expr,
        "invoice_amount_expr": invoice_amount_expr,
        "payable_amount_expr": payable_amount_expr,
        "paid_now_amount_expr": paid_now_amount_expr,
        "balance_amount_expr": balance_amount_expr,
        "already_paid_amount_expr": already_paid_amount_expr,
        "payment_amount_expr": payment_amount_expr,
        "tds_amount_expr": tds_amount_expr,
        "net_paid_amount_expr": net_paid_amount_expr,
        "patient_expr": patient_expr,
        "passport_expr": passport_expr,
        "hospital_expr": hospital_expr,
    }


def _build_agent_payment_where(filters):
    meta = _agent_payment_meta()
    expr = _agent_payment_base_exprs(meta)

    where = ["1=1"]
    params = []

    if filters.get("from_date"):
        where.append(f"{expr['payment_date_expr']} >= %s")
        params.append(filters.get("from_date"))

    if filters.get("to_date"):
        where.append(f"{expr['payment_date_expr']} <= %s")
        params.append(filters.get("to_date"))

    if filters.get("agent"):
        where.append(f"""
            (
                {expr['agent_code_expr']} LIKE %s
                OR {expr['agent_name_expr']} LIKE %s
            )
        """)
        like = f"%{filters.get('agent')}%"
        params.extend([like, like])

    if filters.get("invoice_number"):
        where.append(f"{expr['invoice_number_expr']} LIKE %s")
        params.append(f"%{filters.get('invoice_number')}%")

    if filters.get("payment_no"):
        where.append(f"{expr['payment_no_expr']} LIKE %s")
        params.append(f"%{filters.get('payment_no')}%")

    if filters.get("payment_mode") and filters.get("payment_mode").upper() != "ALL":
        where.append(f"UPPER(COALESCE({expr['payment_mode_expr']},'')) = %s")
        params.append(filters.get("payment_mode").upper())

    if filters.get("payment_status") and filters.get("payment_status").upper() != "ALL":
        where.append(f"UPPER(COALESCE({expr['payment_status_expr']},'')) = %s")
        params.append(filters.get("payment_status").upper())

    if filters.get("currency_code") and filters.get("currency_code").upper() != "ALL":
        where.append(f"UPPER(COALESCE({expr['currency_expr']},'')) = %s")
        params.append(filters.get("currency_code").upper())

    search = filters.get("search") or filters.get("like")
    if search:
        like = f"%{search}%"
        where.append(f"""
            (
                {expr['payment_no_expr']} LIKE %s
                OR {expr['invoice_number_expr']} LIKE %s
                OR {expr['agent_code_expr']} LIKE %s
                OR {expr['agent_name_expr']} LIKE %s
                OR {expr['patient_expr']} LIKE %s
                OR {expr['passport_expr']} LIKE %s
                OR {expr['utr_expr']} LIKE %s
                OR {expr['bank_ref_expr']} LIKE %s
            )
        """)
        params.extend([like, like, like, like, like, like, like, like])

    return " AND ".join(where), params, meta, expr


def _agent_payment_from_clause(meta):
    return f"""
        FROM agent_payment_detail apd
        {_agent_payment_join_clause(meta)}
        {_agent_payment_invoice_join_clause(meta)}
    """


def get_agent_payment_report_dashboard(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            COUNT(DISTINCT {expr['payment_no_expr']}) AS total_payments,
            COUNT(DISTINCT {expr['invoice_number_expr']}) AS invoices_paid,
            COUNT(DISTINCT {expr['agent_code_expr']}) AS unique_agents_paid,
            COALESCE(SUM({expr['payable_amount_expr']}),0) AS total_agency_payable,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) AS total_paid_amount,
            COALESCE(SUM({expr['tds_amount_expr']}),0) AS total_tds_amount,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) - COALESCE(SUM({expr['tds_amount_expr']}),0) AS net_transfer_amount,
            COALESCE(SUM({expr['balance_amount_expr']}),0) AS outstanding_amount
        {from_clause}
        WHERE {where}
    """

    filters_data = {
        "agents": [],
        "payment_modes": [],
        "payment_statuses": [],
        "currencies": [],
    }

    if meta["agent_code_col"]:
        filters_data["agents"] = [r["value"] for r in _fetch_all(f"""
            SELECT DISTINCT COALESCE({meta["agent_name_col"] if meta["agent_name_col"] else meta["agent_code_col"]}, {meta["agent_code_col"]}) AS value
            FROM agent_payment_master
            WHERE {meta["agent_code_col"]} IS NOT NULL
              AND TRIM({meta["agent_code_col"]}) <> ''
            ORDER BY value
        """)]

    if meta["payment_mode_col"]:
        filters_data["payment_modes"] = [r["value"] for r in _fetch_all(f"""
            SELECT DISTINCT {meta["payment_mode_col"]} AS value
            FROM agent_payment_master
            WHERE {meta["payment_mode_col"]} IS NOT NULL
              AND TRIM({meta["payment_mode_col"]}) <> ''
            ORDER BY {meta["payment_mode_col"]}
        """)]

    if meta["payment_status_col"]:
        filters_data["payment_statuses"] = [r["value"] for r in _fetch_all(f"""
            SELECT DISTINCT {meta["payment_status_col"]} AS value
            FROM agent_payment_master
            WHERE {meta["payment_status_col"]} IS NOT NULL
              AND TRIM({meta["payment_status_col"]}) <> ''
            ORDER BY {meta["payment_status_col"]}
        """)]

    if meta["currency_code_col"]:
        filters_data["currencies"] = [r["value"] for r in _fetch_all(f"""
            SELECT DISTINCT {meta["currency_code_col"]} AS value
            FROM agent_payment_master
            WHERE {meta["currency_code_col"]} IS NOT NULL
              AND TRIM({meta["currency_code_col"]}) <> ''
            ORDER BY {meta["currency_code_col"]}
        """)]

    return {
        "summary": _fetch_one(sql, params),
        "filters": filters_data,
    }


def get_agent_payment_report_rows(report_type, filters):
    report_type = (report_type or "payment_detail").strip().lower()

    if report_type == "agent_outstanding":
        return _agent_payment_outstanding_report(filters)

    if report_type == "agent_summary":
        return _agent_payment_agent_summary(filters)

    if report_type == "monthly_payment":
        return _agent_payment_monthly_report(filters)

    if report_type == "tds_register":
        return _agent_payment_tds_register(filters)

    return _agent_payment_detail_report(filters)


def _agent_payment_detail_report(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            {expr['payment_no_expr']} AS payment_no,
            {expr['payment_date_select']} AS payment_date,
            {expr['agent_code_expr']} AS agent_code,
            {expr['agent_name_expr']} AS agent_name,
            {expr['invoice_number_expr']} AS invoice_number,
            {expr['patient_expr']} AS patient_case,
            {expr['passport_expr']} AS patient_passport,
            {expr['hospital_expr']} AS hospital_name,
            {expr['currency_expr']} AS currency,
            {expr['invoice_amount_expr']} AS invoice_amount,
            {expr['payable_amount_expr']} AS agency_pay_amount,
            {expr['already_paid_amount_expr']} AS already_paid_amount,
            {expr['paid_now_amount_expr']} AS paid_amount,
            {expr['balance_amount_expr']} AS balance_amount,
            {expr['tds_amount_expr']} AS tds_amount,
            {expr['net_paid_amount_expr']} AS net_paid_amount,
            {expr['payment_mode_expr']} AS payment_mode,
            {expr['utr_expr']} AS utr_reference,
            {expr['bank_ref_expr']} AS bank_reference,
            {expr['payment_status_expr']} AS payment_status,
            {expr['remarks_expr']} AS remarks,
            {expr['entered_by_expr']} AS entered_by,
            {expr['entered_on_expr']} AS entered_on
        {from_clause}
        WHERE {where}
        ORDER BY {expr['payment_date_expr']} DESC, {expr['payment_no_expr']} DESC, {expr['invoice_number_expr']} DESC
    """

    return _fetch_all(sql, params)


def _agent_payment_outstanding_report(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            {expr['agent_code_expr']} AS agent_code,
            {expr['agent_name_expr']} AS agent_name,
            COUNT(DISTINCT {expr['invoice_number_expr']}) AS invoice_count,
            COALESCE(SUM({expr['invoice_amount_expr']}),0) AS invoice_amount,
            COALESCE(SUM({expr['payable_amount_expr']}),0) AS total_agency_payable,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) AS total_paid_amount,
            COALESCE(SUM({expr['balance_amount_expr']}),0) AS outstanding_amount,
            MAX({expr['payment_date_select']}) AS last_payment_date
        {from_clause}
        WHERE {where}
        GROUP BY {expr['agent_code_expr']}, {expr['agent_name_expr']}
        ORDER BY outstanding_amount DESC, total_agency_payable DESC
    """

    return _fetch_all(sql, params)


def _agent_payment_agent_summary(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            {expr['agent_code_expr']} AS agent_code,
            {expr['agent_name_expr']} AS agent_name,
            COUNT(DISTINCT {expr['payment_no_expr']}) AS payment_count,
            COUNT(DISTINCT {expr['invoice_number_expr']}) AS invoice_count,
            COALESCE(SUM({expr['payable_amount_expr']}),0) AS gross_payable,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) AS gross_paid,
            COALESCE(SUM({expr['tds_amount_expr']}),0) AS tds_amount,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) - COALESCE(SUM({expr['tds_amount_expr']}),0) AS net_paid,
            COALESCE(SUM({expr['balance_amount_expr']}),0) AS balance_amount,
            MAX({expr['payment_date_select']}) AS last_payment_date
        {from_clause}
        WHERE {where}
        GROUP BY {expr['agent_code_expr']}, {expr['agent_name_expr']}
        ORDER BY gross_paid DESC, last_payment_date DESC
    """

    return _fetch_all(sql, params)


def _agent_payment_monthly_report(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            DATE_FORMAT({expr['payment_date_expr']}, '%Y-%m') AS month_name,
            COUNT(DISTINCT {expr['payment_no_expr']}) AS payment_count,
            COUNT(DISTINCT {expr['invoice_number_expr']}) AS invoice_count,
            COUNT(DISTINCT {expr['agent_code_expr']}) AS agent_count,
            COALESCE(SUM({expr['payable_amount_expr']}),0) AS gross_payable,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) AS gross_paid,
            COALESCE(SUM({expr['tds_amount_expr']}),0) AS tds_amount,
            COALESCE(SUM({expr['paid_now_amount_expr']}),0) - COALESCE(SUM({expr['tds_amount_expr']}),0) AS net_paid,
            COALESCE(SUM({expr['balance_amount_expr']}),0) AS balance_amount
        {from_clause}
        WHERE {where}
        GROUP BY DATE_FORMAT({expr['payment_date_expr']}, '%Y-%m')
        ORDER BY month_name DESC
    """

    return _fetch_all(sql, params)


def _agent_payment_tds_register(filters):
    where, params, meta, expr = _build_agent_payment_where(filters)
    from_clause = _agent_payment_from_clause(meta)

    sql = f"""
        SELECT
            {expr['payment_no_expr']} AS payment_no,
            {expr['payment_date_select']} AS payment_date,
            {expr['agent_code_expr']} AS agent_code,
            {expr['agent_name_expr']} AS agent_name,
            {expr['invoice_number_expr']} AS invoice_number,
            {expr['currency_expr']} AS currency,
            {expr['paid_now_amount_expr']} AS gross_paid,
            {expr['tds_amount_expr']} AS tds_amount,
            CASE
                WHEN {expr['paid_now_amount_expr']} = 0 THEN 0
                ELSE ROUND(({expr['tds_amount_expr']} / {expr['paid_now_amount_expr']}) * 100, 2)
            END AS tds_percent,
            {expr['paid_now_amount_expr']} - {expr['tds_amount_expr']} AS net_paid,
            {expr['utr_expr']} AS utr_reference,
            {expr['payment_status_expr']} AS payment_status
        {from_clause}
        WHERE {where}
        ORDER BY {expr['payment_date_expr']} DESC, {expr['payment_no_expr']} DESC
    """

    return _fetch_all(sql, params)

# ============================================================
# COMMON EXCEL / PDF EXPORTS FOR REPORTS
# ============================================================

def _export_report_config(report_key, filters):
    report_key = (report_key or "").strip().lower()

    if report_key == "queue-ageing":
        return {
            "title": "Queue Ageing Report",
            "sheet_name": "Queue Ageing",
            "filename_base": "queue_ageing_report",
            "rows": get_queue_ageing_report(filters),
        }

    if report_key == "invoice-management":
        report_type = filters.get("report_type", "invoice_detail")
        report_titles = {
            "invoice_detail": "Invoice Detail Report",
            "collection_report": "Collection Report",
            "hospital_outstanding": "Hospital Outstanding Report",
            "invoice_ageing": "Invoice Ageing Report",
            "monthly_revenue": "Monthly Revenue Report",
        }

        title = report_titles.get(report_type, "Invoice Management Report")

        return {
            "title": title,
            "sheet_name": "Invoice Report",
            "filename_base": report_type,
            "rows": get_invoice_report_rows(report_type, filters),
        }

    if report_key == "agent-payment":
        report_type = filters.get("report_type", "payment_detail")
        report_titles = {
            "payment_detail": "Agent Payment Detail Report",
            "agent_outstanding": "Agent Outstanding Report",
            "agent_summary": "Agent Wise Payment Summary",
            "monthly_payment": "Monthly Agent Payment Trend",
            "tds_register": "Agent Payment TDS Register",
        }

        title = report_titles.get(report_type, "Agent Payment Report")

        return {
            "title": title,
            "sheet_name": "Agent Payment",
            "filename_base": f"agent_payment_{report_type}",
            "rows": get_agent_payment_report_rows(report_type, filters),
        }

    if report_key == "sla":
        return {
            "title": "SLA Report",
            "sheet_name": "SLA Report",
            "filename_base": "sla_report",
            "rows": get_sla_report_rows(filters),
        }

    raise ValueError("Invalid report key")


def _clean_export_cell_value(value):
    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if value is None:
        return ""

    return value



def _report_as_on_text():
    return f"Report as on {datetime.now().strftime('%d-%m-%Y')}"

def _export_rows_to_excel(title, sheet_name, rows, filename):
    from io import BytesIO
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] if sheet_name else "Report"

    total_cols = 12
    if rows:
        total_cols = max(len(rows[0].keys()), 1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F4E78")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    as_on_cell = ws.cell(row=2, column=1)
    as_on_cell.value = _report_as_on_text()
    as_on_cell.font = Font(size=11, bold=True, color="1F4E78")
    as_on_cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    if not rows:
        ws.cell(row=4, column=1).value = "No data available"
        ws.cell(row=4, column=1).font = Font(bold=True)
    else:
        headers = list(rows[0].keys())

        header_fill = PatternFill("solid", fgColor="305496")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = header.replace("_", " ").title()
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row_num, row in enumerate(rows, 5):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = _clean_export_cell_value(row.get(header))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin_border

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = ws.dimensions

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            # Use a real worksheet cell to get the column letter.
            # This avoids openpyxl MergedCell errors caused by the merged report title row.
            col_letter = ws.cell(row=4, column=col_idx).column_letter

            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))

            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _export_rows_to_pdf(title, rows, filename):
    from io import BytesIO
    from flask import send_file

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except Exception as exc:
        raise Exception("PDF export requires reportlab. Please run: pip install reportlab") from exc

    output = BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=22,
        bottomMargin=18,
    )

    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]
    normal_style.fontSize = 6
    normal_style.leading = 7

    header_style = styles["Normal"].clone("ReportTableHeader")
    header_style.fontName = "Helvetica-Bold"
    header_style.fontSize = 6
    header_style.leading = 7
    header_style.textColor = colors.white
    header_style.alignment = 1

    as_on_style = styles["Normal"].clone("ReportAsOn")
    as_on_style.fontName = "Helvetica-Bold"
    as_on_style.fontSize = 9
    as_on_style.leading = 11
    as_on_style.textColor = colors.HexColor("#1F4E78")
    as_on_style.alignment = 1

    elements = [
        Paragraph(f"<b>{title}</b>", styles["Title"]),
        Paragraph(_report_as_on_text(), as_on_style),
        Spacer(1, 10),
    ]

    if not rows:
        elements.append(Paragraph("No data available", styles["Normal"]))
    else:
        headers = list(rows[0].keys())

        # Keep PDF readable by limiting very wide datasets.
        # Excel export still contains every column.
        max_pdf_cols = 12
        display_headers = headers[:max_pdf_cols]

        data = [[Paragraph(h.replace("_", " ").title(), header_style) for h in display_headers]]

        for row in rows:
            data.append([
                Paragraph(str(_clean_export_cell_value(row.get(h))), normal_style)
                for h in display_headers
            ])

        available_width = landscape(A4)[0] - doc.leftMargin - doc.rightMargin
        col_width = available_width / max(len(display_headers), 1)

        table = Table(
            data,
            repeatRows=1,
            colWidths=[col_width] * len(display_headers),
        )

        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))

        elements.append(table)

        if len(headers) > max_pdf_cols:
            elements.append(Spacer(1, 8))
            elements.append(Paragraph(
                f"Note: PDF includes first {max_pdf_cols} columns for readability. Export Excel includes all columns.",
                styles["Normal"],
            ))

    doc.build(elements)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


def export_report_excel(report_key, filters):
    config = _export_report_config(report_key, filters or {})
    return _export_rows_to_excel(
        title=config["title"],
        sheet_name=config["sheet_name"],
        rows=config["rows"],
        filename=f'{config["filename_base"]}.xlsx',
    )


def export_report_pdf(report_key, filters):
    config = _export_report_config(report_key, filters or {})
    return _export_rows_to_pdf(
        title=config["title"],
        rows=config["rows"],
        filename=f'{config["filename_base"]}.pdf',
    )

