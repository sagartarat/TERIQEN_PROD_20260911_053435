import json
import os
from datetime import datetime
from io import BytesIO

from flask import render_template, request, jsonify, current_app, send_file
from flask_login import login_required, current_user

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.common.db import get_db
from app.common.file_helpers import save_uploaded_files
from .dashboard_service import get_hospital_dashboard_context, get_hospital_dashboard_export_context


BRAND_NAVY = "163A70"
BRAND_BLUE = "1E4F95"
BRAND_LIGHT = "EAF1FB"
GRID = "D9E4F1"
TEXT = "15345F"
MUTED = "6D7F99"
GREEN = "DFF6EA"
GREEN_TEXT = "147A3D"
AMBER = "FFF1D6"
AMBER_TEXT = "9A5B00"
RED = "FFE5E5"
RED_TEXT = "A32020"
PURPLE = "EFE7FF"
PURPLE_TEXT = "5B2DB7"
WHITE = "FFFFFF"
GREY = "F6F9FD"


def generate_stub_message_id():
    return f"WA_STUB_{datetime.now().strftime('%Y%m%d%H%M%S%f')[:17]}"


def _excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (int, float, bool, datetime)):
        return value
    return str(value)


def _money(value):
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _int(value):
    try:
        return int(float(value or 0))
    except Exception:
        return 0


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _font(color=TEXT, bold=False, size=10):
    return Font(name="Calibri", color=color, bold=bold, size=size)


def _thin_border():
    side = Side(style="thin", color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


def _section_title(ws, row, start_col, end_col, title):
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row=row, column=start_col)
    cell.value = title
    cell.fill = _fill(BRAND_LIGHT)
    cell.font = _font(BRAND_NAVY, True, 12)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = _thin_border()
    ws.row_dimensions[row].height = 24


def _set_col_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _write_kpi_card(ws, row, col, label, value, fill_color, value_format=None):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 1)

    label_cell = ws.cell(row=row, column=col)
    label_cell.value = label
    label_cell.fill = _fill(fill_color)
    label_cell.font = _font(WHITE, True, 9)
    label_cell.alignment = Alignment(horizontal="center", vertical="center")

    value_cell = ws.cell(row=row + 1, column=col)
    value_cell.value = value
    value_cell.fill = _fill("FFFFFF")
    value_cell.font = _font(TEXT, True, 17)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    if value_format:
        value_cell.number_format = value_format

    for r in range(row, row + 3):
        for c in range(col, col + 2):
            ws.cell(row=r, column=c).border = _thin_border()
            if r > row:
                ws.cell(row=r, column=c).fill = _fill("FFFFFF")

    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 22
    ws.row_dimensions[row + 2].height = 18


def _style_data_table(ws, start_row, headers, rows, table_name):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx)
        cell.value = header
        cell.fill = _fill(BRAND_NAVY)
        cell.font = _font(WHITE, True, 10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()

    data_start = start_row + 1
    if rows:
        for row_idx, row_data in enumerate(rows, data_start):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = _excel_value(value)
                cell.border = _thin_border()
                cell.font = _font(TEXT, False, 10)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_idx % 2 == 0:
                    cell.fill = _fill(GREY)
    else:
        ws.cell(row=data_start, column=1).value = "No records available"
        ws.cell(row=data_start, column=1).font = _font(MUTED, True, 10)
        ws.cell(row=data_start, column=1).border = _thin_border()

    end_row = max(data_start, start_row + len(rows))
    end_col = len(headers)
    table_ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    tab = Table(displayName=table_name, ref=table_ref)
    style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    try:
        ws.add_table(tab)
    except Exception:
        pass

    ws.auto_filter.ref = table_ref
    ws.freeze_panes = f"A{data_start}"

    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.border = _thin_border()

    return end_row


def _add_status_conditional_formatting(ws, status_col_letter, first_row, last_row):
    if last_row < first_row:
        return
    rng = f"{status_col_letter}{first_row}:{status_col_letter}{last_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"BREACHED"'], fill=_fill(RED), font=_font(RED_TEXT, True, 10)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NEAR BREACH"'], fill=_fill(AMBER), font=_font(AMBER_TEXT, True, 10)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ESCALATED"'], fill=_fill(PURPLE), font=_font(PURPLE_TEXT, True, 10)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"MATCHED"'], fill=_fill(GREEN), font=_font(GREEN_TEXT, True, 10)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PARTIAL MATCH"'], fill=_fill(AMBER), font=_font(AMBER_TEXT, True, 10)))


def _prepare_print(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45


def _build_erp_dashboard_excel(context):
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    summary = context.get("summary") or {}
    finance = context.get("finance") or {}
    date_range = context.get("selected_date_range") or "Today"
    exported_on = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    ws = wb.create_sheet("Executive Dashboard")
    _prepare_print(ws, landscape=True)
    _set_col_widths(ws, {
        "A": 3, "B": 18, "C": 14, "D": 18, "E": 14, "F": 18, "G": 14,
        "H": 18, "I": 14, "J": 18, "K": 14, "L": 18, "M": 14
    })

    ws.merge_cells("B2:M3")
    title = ws["B2"]
    title.value = "Hospital ERP Dashboard Export"
    title.fill = _fill(BRAND_NAVY)
    title.font = _font(WHITE, True, 20)
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 10

    ws["B4"] = "Date Range"
    ws["C4"] = date_range
    ws["E4"] = "Exported On"
    ws["F4"] = exported_on
    ws["H4"] = "Report Type"
    ws["I4"] = "ERP Management Pack"
    for cell in ["B4", "E4", "H4"]:
        ws[cell].font = _font(MUTED, True, 10)
    for cell in ["C4", "F4", "I4"]:
        ws[cell].font = _font(TEXT, True, 10)

    _section_title(ws, 6, 2, 13, "Operational SLA Snapshot")
    _write_kpi_card(ws, 8, 2, "ASSIGNED", _int(summary.get("assigned")), BRAND_BLUE)
    _write_kpi_card(ws, 8, 4, "IN PROGRESS", _int(summary.get("in_progress")), "15803D")
    _write_kpi_card(ws, 8, 6, "CLOSED", _int(summary.get("closed_today")), "15803D")
    _write_kpi_card(ws, 8, 8, "NEAR BREACH", _int(summary.get("near_breach")), "D97706")
    _write_kpi_card(ws, 8, 10, "BREACHED", _int(summary.get("breached")), "B91C1C")
    _write_kpi_card(ws, 8, 12, "ESCALATED", _int(summary.get("escalated")), "6D28D9")

    _section_title(ws, 13, 2, 13, "Invoice & Finance Snapshot")
    _write_kpi_card(ws, 15, 2, "INVOICES", _int(finance.get("invoices_generated")), BRAND_BLUE)
    _write_kpi_card(ws, 15, 4, "PENDING", _int(finance.get("pending_invoices")), "D97706")
    _write_kpi_card(ws, 15, 6, "PAID", _int(finance.get("paid_invoices")), "15803D")
    _write_kpi_card(ws, 15, 8, "PARTIAL", _int(finance.get("partially_paid")), "D97706")
    _write_kpi_card(ws, 15, 10, "OUTSTANDING", _money(finance.get("outstanding_amount")), "B91C1C", '#,##0.00')
    _write_kpi_card(ws, 15, 12, "PAYMENT RATIO", _money(finance.get("payment_ratio")) / 100, BRAND_BLUE, '0%')

    _section_title(ws, 20, 2, 6, "SLA Health Source Data")
    sla_rows = [
        ["Within SLA", _int(summary.get("within_sla"))],
        ["Near Breach", _int(summary.get("near_breach"))],
        ["Breached", _int(summary.get("breached"))],
        ["Escalated", _int(summary.get("escalated"))],
    ]
    for idx, row_data in enumerate(sla_rows, 22):
        ws.cell(idx, 2).value = row_data[0]
        ws.cell(idx, 3).value = row_data[1]
        for c in range(2, 4):
            ws.cell(idx, c).border = _thin_border()
            ws.cell(idx, c).font = _font(TEXT, True if c == 2 else False, 10)
            ws.cell(idx, c).fill = _fill("FFFFFF" if idx % 2 else GREY)
    ws["B21"] = "Status"
    ws["C21"] = "Count"
    for c in ["B21", "C21"]:
        ws[c].fill = _fill(BRAND_NAVY)
        ws[c].font = _font(WHITE, True, 10)
        ws[c].alignment = Alignment(horizontal="center")
        ws[c].border = _thin_border()

    chart = DoughnutChart()
    labels = Reference(ws, min_col=2, min_row=22, max_row=25)
    data = Reference(ws, min_col=3, min_row=21, max_row=25)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.title = "SLA Health Mix"
    chart.holeSize = 55
    ws.add_chart(chart, "E21")

    _section_title(ws, 20, 9, 13, "Finance Mix Source Data")
    finance_rows = [
        ["Paid", _int(finance.get("paid_invoices"))],
        ["Pending", _int(finance.get("pending_invoices"))],
        ["Partial", _int(finance.get("partially_paid"))],
        ["Collections", _money(finance.get("collections_today"))],
    ]
    ws["I21"] = "Metric"
    ws["J21"] = "Value"
    for c in ["I21", "J21"]:
        ws[c].fill = _fill(BRAND_NAVY)
        ws[c].font = _font(WHITE, True, 10)
        ws[c].alignment = Alignment(horizontal="center")
        ws[c].border = _thin_border()
    for idx, row_data in enumerate(finance_rows, 22):
        ws.cell(idx, 9).value = row_data[0]
        ws.cell(idx, 10).value = row_data[1]
        for c in range(9, 11):
            ws.cell(idx, c).border = _thin_border()
            ws.cell(idx, c).font = _font(TEXT, True if c == 9 else False, 10)
            ws.cell(idx, c).fill = _fill("FFFFFF" if idx % 2 else GREY)
    bar = BarChart()
    bar.type = "col"
    bar.title = "Finance Collection View"
    bar.y_axis.title = "Value"
    bar.x_axis.title = "Metric"
    data = Reference(ws, min_col=10, min_row=21, max_row=25)
    cats = Reference(ws, min_col=9, min_row=22, max_row=25)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.height = 7
    bar.width = 12
    ws.add_chart(bar, "K21")

    ws.freeze_panes = "B6"

    ws = wb.create_sheet("Operation Queue")
    _prepare_print(ws, landscape=True)
    operation_rows = []
    for row in context.get("operation_queue_rows") or []:
        operation_rows.append([
            row.get("channel_thread_id"),
            row.get("patient_name") or row.get("patgiven_nm"),
            row.get("patpassport_no"),
            row.get("hospital_shrtnm"),
            row.get("wrkflw_status"),
            row.get("priority_status"),
            row.get("assigned_to"),
            row.get("sla_state"),
        ])
    headers = ["Thread ID", "Patient Name", "Passport No", "Hospital", "Workflow Status", "Priority", "Assigned To", "SLA Status"]
    _section_title(ws, 1, 1, len(headers), "Live Operation Queue - Full Export")
    end_row = _style_data_table(ws, 3, headers, operation_rows, "OperationQueueTable")
    _set_col_widths(ws, {"A": 30, "B": 24, "C": 18, "D": 20, "E": 18, "F": 13, "G": 18, "H": 16})
    _add_status_conditional_formatting(ws, "H", 4, end_row)

    ws = wb.create_sheet("Outstanding Invoices")
    _prepare_print(ws, landscape=True)
    invoice_rows = []
    for row in context.get("invoice_rows") or []:
        invoice_rows.append([
            row.get("invoice_no"),
            row.get("patient_name"),
            row.get("hospital_name") or row.get("hospital_shrtnm"),
            _money(row.get("invoice_value")),
            _money(row.get("paid_amount")),
            _money(row.get("balance_amount")),
            row.get("invoice_status"),
        ])
    headers = ["Invoice No", "Patient Name", "Hospital", "Invoice Value", "Paid Amount", "Balance", "Latest Status"]
    _section_title(ws, 1, 1, len(headers), "Invoice Management Summary - Outstanding Invoices Only")
    end_row = _style_data_table(ws, 3, headers, invoice_rows, "OutstandingInvoiceTable")
    _set_col_widths(ws, {"A": 42, "B": 24, "C": 30, "D": 16, "E": 16, "F": 16, "G": 18})
    for col in ["D", "E", "F"]:
        for r in range(4, end_row + 1):
            ws[f"{col}{r}"].number_format = '#,##0.00'
    _add_status_conditional_formatting(ws, "G", 4, end_row)

    ws = wb.create_sheet("Alerts")
    _prepare_print(ws, landscape=True)
    alert_rows = []
    for alert in context.get("quick_alerts") or []:
        alert_rows.append([alert.get("title"), alert.get("description"), alert.get("badge_label")])
    headers = ["Alert", "Description", "Category"]
    _section_title(ws, 1, 1, len(headers), "Quick Alerts & Exceptions")
    end_row = _style_data_table(ws, 3, headers, alert_rows, "AlertTable")
    _set_col_widths(ws, {"A": 26, "B": 80, "C": 18})

    ws = wb.create_sheet("Chart Data")
    _prepare_print(ws, landscape=True)
    chart_rows = []
    for item in context.get("operation_status_chart") or []:
        chart_rows.append(["Operation Status", item.get("label"), item.get("value")])
    for item in context.get("finance_trend_chart") or []:
        chart_rows.append(["Finance Trend", item.get("label"), item.get("value")])
    headers = ["Section", "Label", "Value"]
    _section_title(ws, 1, 1, len(headers), "Dashboard Chart Source Data")
    _style_data_table(ws, 3, headers, chart_rows, "ChartDataTable")
    _set_col_widths(ws, {"A": 24, "B": 28, "C": 16})

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 90

    wb.active = 0
    return wb


def register_dashboard_routes(app):
    @app.route("/dashboard")
    @login_required
    def dashboard():
        conn = None
        try:
            conn = get_db()
            context = get_hospital_dashboard_context(conn, request.args)
            context["user"] = current_user
            return render_template("hospitaldashboard.html", **context)
        except Exception as e:
            print("dashboard route error:", e)
            return render_template(
                "hospitaldashboard.html",
                user=current_user,
                summary={},
                finance={},
                operation_queue_rows=[],
                invoice_rows=[],
                quick_alerts=[],
                operation_status_chart=[],
                finance_trend_chart=[],
                dashboard_last_updated="",
                selected_hospital="",
                selected_department="",
                selected_date_range="",
                now=datetime.now().strftime("%d-%b-%Y %I:%M %p"),
                operation_queue_pagination={},
                invoice_pagination={}
            )
        finally:
            try:
                if conn:
                    conn.close()
            except Exception:
                pass

    @app.route("/dashboard/export-excel")
    @login_required
    def dashboard_export_excel():
        conn = None
        try:
            conn = get_db()
            context = get_hospital_dashboard_export_context(conn, request.args)
            wb = _build_erp_dashboard_excel(context)

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            file_name = f"Hospital_ERP_Dashboard_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return send_file(
                output,
                as_attachment=True,
                download_name=file_name,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            print("dashboard_export_excel route error:", e)
            return jsonify({
                "success": False,
                "message": f"Unable to export dashboard Excel: {str(e)}"
            }), 500
        finally:
            try:
                if conn:
                    conn.close()
            except Exception:
                pass

    @app.route("/whatsapp-stub")
    @login_required
    def whatsapp_stub():
        return render_template("whatsapp_stub.html", user=current_user)

    @app.route("/whatsapp-stub/push", methods=["POST"])
    @login_required
    def whatsapp_stub_push():
        conn = None
        cursor = None

        try:
            mobile_number = (request.form.get("mobile_number") or "").strip()
            message = (request.form.get("message") or "").strip()
            uploaded_files = request.files.getlist("attachments")

            if not message:
                return jsonify({"success": False, "message": "Message is required."}), 400

            channel_message_id = generate_stub_message_id()

            base_upload_folder = os.path.join(current_app.root_path, "static", "uploads", "whatsapp_docs")

            file_result = save_uploaded_files(
                files=uploaded_files,
                base_folder=base_upload_folder,
                channel_message_id=channel_message_id,
            )

            if not file_result.get("success"):
                return jsonify({"success": False, "message": file_result.get("message") or "File upload failed."}), 400

            attachment_count = file_result.get("attachment_count", 0)
            attachment_json = file_result.get("attachment_json", "[]")

            payload_json = json.dumps(
                {
                    "source": "whatsapp_stub",
                    "channel_code": "WHATSAPP",
                    "channel_message_id": channel_message_id,
                    "sender_mobile": mobile_number,
                    "raw_message_text": message,
                    "attachment_count": attachment_count,
                    "attachments": file_result.get("attachments", []),
                },
                ensure_ascii=False
            )

            conn = get_db()
            cursor = conn.cursor()

            insert_sql = """
                INSERT INTO channel_message_log
                (
                    channel_code,
                    channel_message_id,
                    sender_id,
                    sender_name,
                    sender_mobile,
                    raw_message_text,
                    raw_payload_json,
                    attachment_count,
                    attachment_json,
                    received_on,
                    parsed_flag,
                    parse_status,
                    spam_flag,
                    junk_flag,
                    duplicate_flag,
                    remarks
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    NOW(),
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """

            cursor.execute(
                insert_sql,
                (
                    "WHATSAPP",
                    channel_message_id,
                    mobile_number if mobile_number else None,
                    getattr(current_user, "username", None) or getattr(current_user, "user_code", None) or "Stub User",
                    mobile_number if mobile_number else None,
                    message,
                    payload_json,
                    attachment_count,
                    attachment_json,
                    0,
                    "PENDING",
                    0,
                    0,
                    0,
                    "Inserted from Whatsapp Stub"
                )
            )

            channel_message_log_id = cursor.lastrowid
            conn.commit()

            return jsonify({
                "success": True,
                "message": "Whatsapp stub message submitted successfully.",
                "channel_message_log_id": channel_message_log_id,
                "channel_message_id": channel_message_id,
                "attachment_count": attachment_count
            })

        except Exception as e:
            try:
                if conn:
                    conn.rollback()
            except Exception:
                pass

            print("whatsapp_stub_push error:", e)
            return jsonify({"success": False, "message": f"Failed to save WhatsApp stub message: {str(e)}"}), 500

        finally:
            try:
                if cursor:
                    cursor.close()
            except Exception:
                pass

            try:
                if conn:
                    conn.close()
            except Exception:
                pass
